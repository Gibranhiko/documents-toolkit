"""Extrae fecha y monto de comprobantes SPEI (PDF) y los concentra en un Excel.

Opcionalmente reparte el monto en dos porciones configurables (--split / etiquetas),
por ejemplo para dividir una comisión.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from doctk.common.paths import expand_inputs

LABEL_FECHA = r"^Fecha de operación$"
LABEL_IMPORTE = r"^Importe$"
CURR_RE = re.compile(r"\$\s?\d{1,3}(?:,\d{3})*(?:\.\d{2})")
FECHA_RE = re.compile(r"\b\d{2}/\d{2}/\d{4}\s+\d{1,2}:\d{2}\s*(?:AM|PM)?\b", re.IGNORECASE)


def add_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument(
        "-i", "--input", required=True, nargs="+",
        help="PDFs, carpetas o patrones glob con comprobantes SPEI.",
    )
    parser.add_argument(
        "-o", "--output", type=Path, default=Path("resultado_fechas_montos.xlsx"),
        help="Excel de salida (por defecto: resultado_fechas_montos.xlsx).",
    )
    parser.add_argument(
        "--split", type=float, default=None, metavar="FRACCION",
        help="Fracción (0-1) para dividir el monto en dos columnas (ej. 0.10).",
    )
    parser.add_argument(
        "--label-a", default="ParteA",
        help="Encabezado de la porción --split (por defecto: ParteA).",
    )
    parser.add_argument(
        "--label-b", default="ParteB",
        help="Encabezado del resto tras --split (por defecto: ParteB).",
    )


def _to_float(money: str) -> float:
    s = str(money).replace("$", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def _fecha_importe_from_page(page) -> tuple[str, str]:
    fecha, importe = "", ""
    tables = page.extract_tables({
        "vertical_strategy": "lines",
        "horizontal_strategy": "lines",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "edge_min_length": 20,
    }) or []
    if not tables:
        tables = page.extract_tables({
            "vertical_strategy": "text",
            "horizontal_strategy": "text",
            "text_x_tolerance": 2,
            "text_y_tolerance": 2,
        }) or []
    for tbl in tables:
        for row in tbl or []:
            if not row or len(row) < 2:
                continue
            key = (row[0] or "").replace(" ", " ").strip()
            val = (row[1] or "").replace(" ", " ").strip()
            if not fecha and re.search(LABEL_FECHA, key, re.IGNORECASE):
                m = FECHA_RE.search(val)
                fecha = m.group(0) if m else val
            if not importe and re.search(LABEL_IMPORTE, key, re.IGNORECASE):
                m = CURR_RE.search(val)
                importe = m.group(0) if m else val
    if not fecha or not importe:
        text = (page.extract_text() or "").replace(" ", " ")
        if not fecha:
            m = re.search(r"Fecha de operación\s+([^\n]+)", text, re.IGNORECASE)
            if m:
                cand = m.group(1).strip()
                m2 = FECHA_RE.search(cand)
                fecha = m2.group(0) if m2 else cand
        if not importe:
            m = re.search(r"Importe\s+([^\n]+)", text, re.IGNORECASE)
            if m:
                m2 = CURR_RE.search(m.group(1).strip())
                importe = m2.group(0) if m2 else ""
    return fecha, importe


def _extract_from_pdf(path: Path) -> tuple[str, str]:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            fecha, importe = _fecha_importe_from_page(page)
            if fecha and importe:
                return fecha, importe
    return "", ""


def run(args: argparse.Namespace) -> int:
    import pandas as pd
    from openpyxl.styles import Font

    files = expand_inputs(args.input, {".pdf"})
    if not files:
        print("No se encontraron PDFs en las rutas indicadas.")
        return 1

    print(f"📄 Archivos detectados ({len(files)}):")
    for f in files:
        print("   -", f.name)
    print()

    split = args.split
    if split is not None and not 0 <= split <= 1:
        raise SystemExit("--split debe estar entre 0 y 1.")

    rows = []
    for idx, f in enumerate(files, start=1):
        print(f"[{idx}/{len(files)}] Procesando {f.name} ...", end=" ")
        fecha, importe_txt = _extract_from_pdf(f)
        if not fecha or not importe_txt:
            print("⚠️  no se encontró fecha/importe")
        else:
            print(f"OK → Fecha: {fecha} | Importe: {importe_txt}")
        monto = _to_float(importe_txt) if importe_txt else 0.0
        row = {"Fecha": fecha, "Monto": monto}
        if split is not None:
            parte_a = round(monto * split, 2)
            row[args.label_a] = parte_a
            row[args.label_b] = round(monto - parte_a, 2)
        rows.append(row)

    columns = ["Fecha", "Monto"] + ([args.label_a, args.label_b] if split is not None else [])
    df = pd.DataFrame(rows, columns=columns)

    numeric_cols = [c for c in columns if c != "Fecha"]
    totals = {c: round(df[c].sum(), 2) for c in numeric_cols}

    print("\n✅ Totales:")
    for c, v in totals.items():
        print(f"   {c}: {v}")

    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Hoja1")
        ws = writer.sheets["Hoja1"]
        last = ws.max_row + 1
        ws.cell(row=last, column=1, value="TOTAL")
        for col_idx, c in enumerate(numeric_cols, start=2):
            ws.cell(row=last, column=col_idx, value=totals[c])
        for col_idx in range(1, len(columns) + 1):
            ws.cell(row=last, column=col_idx).font = Font(bold=True)

    print(f"\n📦 Archivo generado: {args.output}")
    return 0
